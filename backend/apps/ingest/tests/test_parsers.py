import io

import openpyxl
from docx import Document
from pptx import Presentation
from pypdf import PdfWriter

from apps.ingest.pipeline.parsers import PARSER_REGISTRY, UnsupportedContentTypeError, get_parser


def test_text_parser():
    parser = get_parser("text/plain")
    parsed = parser.parse(b"hello world", "text/plain")
    assert len(parsed.segments) == 1
    assert parsed.segments[0].type == "prose"
    assert parsed.segments[0].text == "hello world"


def test_markdown_uses_text_parser():
    assert PARSER_REGISTRY["text/markdown"] is PARSER_REGISTRY["text/plain"]


def test_csv_parser():
    payload = b"name,age\nAda,30\nGrace,85\n"
    parsed = get_parser("text/csv").parse(payload, "text/csv")
    assert len(parsed.segments) == 1
    segment = parsed.segments[0]
    assert segment.type == "tabular"
    assert segment.header == "name\tage"
    assert segment.row_range == (1, 2)
    assert "Ada\t30" in segment.text
    assert "Grace\t85" in segment.text


def test_docx_parser_prose_and_table():
    document = Document()
    document.add_paragraph("Some prose text.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Header A"
    table.cell(0, 1).text = "Header B"
    table.cell(1, 0).text = "Row A"
    table.cell(1, 1).text = "Row B"
    buf = io.BytesIO()
    document.save(buf)

    parsed = get_parser(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ).parse(
        buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

    prose = [s for s in parsed.segments if s.type == "prose"]
    tabular = [s for s in parsed.segments if s.type == "tabular"]
    assert "Some prose text." in prose[0].text
    assert tabular[0].header == "Header A\tHeader B"
    assert "Row A\tRow B" in tabular[0].text


def test_xlsx_parser():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["name", "age"])
    sheet.append(["Ada", 30])
    sheet.append(["Grace", 85])
    buf = io.BytesIO()
    workbook.save(buf)

    parsed = get_parser("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet").parse(
        buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    assert len(parsed.segments) == 1
    segment = parsed.segments[0]
    assert segment.header == "name\tage"
    assert segment.row_range == (1, 2)
    assert "Ada\t30" in segment.text


def test_pptx_parser():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    box = slide.shapes.add_textbox(0, 0, 100, 100)
    box.text_frame.text = "Slide one text"
    buf = io.BytesIO()
    presentation.save(buf)

    parsed = get_parser(
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    ).parse(
        buf.getvalue(),
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )

    assert len(parsed.segments) == 1
    assert parsed.segments[0].page_idx == 0
    assert "Slide one text" in parsed.segments[0].text


def test_pdf_parser():
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)

    parsed = get_parser("application/pdf").parse(buf.getvalue(), "application/pdf")

    assert len(parsed.segments) == 1
    assert parsed.segments[0].page_idx == 0
    # A blank page has no extractable text -- an empty string, not a crash.
    assert parsed.segments[0].text == ""


def test_unsupported_content_type_raises():
    try:
        get_parser("image/png")
    except UnsupportedContentTypeError as e:
        assert e.content_type == "image/png"
    else:
        raise AssertionError("expected UnsupportedContentTypeError")
