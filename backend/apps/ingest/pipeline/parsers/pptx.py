import io

from pptx import Presentation

from apps.ingest.pipeline.parsed_document import ParsedDocument, Segment


class PptxParser:
    def parse(self, payload: bytes, content_type: str) -> ParsedDocument:
        presentation = Presentation(io.BytesIO(payload))
        segments = []
        for i, slide in enumerate(presentation.slides):
            texts = [
                shape.text_frame.text
                for shape in slide.shapes
                if shape.has_text_frame and shape.text_frame.text.strip()
            ]
            segments.append(Segment(type="prose", page_idx=i, text="\n".join(texts)))
        return ParsedDocument(segments=segments, source_content_type=content_type)
