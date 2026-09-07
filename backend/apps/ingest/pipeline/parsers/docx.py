import io

from docx import Document

from apps.ingest.pipeline.parsed_document import ParsedDocument, Segment


class DocxParser:
    def parse(self, payload: bytes, content_type: str) -> ParsedDocument:
        document = Document(io.BytesIO(payload))
        segments = []

        prose_text = "\n".join(p.text for p in document.paragraphs)
        if prose_text.strip():
            segments.append(Segment(type="prose", text=prose_text))

        for table in document.tables:
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            if not rows:
                continue
            header, body = rows[0], rows[1:]
            segments.append(
                Segment(
                    type="tabular",
                    text="\n".join("\t".join(row) for row in body),
                    row_range=(1, len(body)),
                    header="\t".join(header),
                )
            )

        return ParsedDocument(segments=segments, source_content_type=content_type)
