import io

import openpyxl

from apps.ingest.pipeline.parsed_document import ParsedDocument, Segment


class XlsxParser:
    def parse(self, payload: bytes, content_type: str) -> ParsedDocument:
        workbook = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        segments = []
        for sheet in workbook.worksheets:
            rows = [
                ["" if cell is None else str(cell) for cell in row]
                for row in sheet.iter_rows(values_only=True)
            ]
            if not rows:
                continue
            header, body = rows[0], rows[1:]
            if not body:
                continue
            segments.append(
                Segment(
                    type="tabular",
                    text="\n".join("\t".join(row) for row in body),
                    row_range=(1, len(body)),
                    header="\t".join(header),
                )
            )
        return ParsedDocument(segments=segments, source_content_type=content_type)
